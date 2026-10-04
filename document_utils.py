"""
document_utils.py

Helpers for turning an uploaded document (PDF / Word / text) into clean,
clause-sized pieces and scoring each piece with the risk classifier.

Pipeline (brief step 2 and 6):
  file bytes -> extract text -> remove page numbers / odd formatting
             -> split into clauses -> predict risk, confidence, action
"""

import io
import re
import unicodedata

import numpy as np
import pandas as pd

CONFIDENCE_THRESHOLD = 0.60
MIN_CHARS, MAX_CHARS = 25, 1000

ACTION_ICONS = {
    "Escalate immediately": "🚨",
    "Flag for review": "⚠️",
    "Auto-clear": "✅",
}
ACTION_ORDER = {"Escalate immediately": 0, "Flag for review": 1, "Auto-clear": 2}


# ---------------------------------------------------------------- policy ----
def review_action(risk, confidence):
    """Tiered review policy (same rule as train_risk_classifier.py)."""
    if confidence < CONFIDENCE_THRESHOLD and risk == "high":
        return "Escalate immediately"
    if confidence < CONFIDENCE_THRESHOLD or risk == "high":
        return "Flag for review"
    return "Auto-clear"


# ------------------------------------------------------- text extraction ----
def extract_text(file_bytes, filename):
    """Return raw text from a .pdf, .docx or .txt file."""
    name = filename.lower()
    if name.endswith(".pdf"):
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(file_bytes))
        pages = [(page.extract_text() or "").splitlines() for page in reader.pages]
        return "\n\n".join("\n".join(p) for p in strip_repeated_lines(pages))
    if name.endswith(".docx"):
        import docx
        document = docx.Document(io.BytesIO(file_bytes))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n\n".join(parts)
    if name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")
    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


def _shape(line):
    """Line with digits masked, so '1/9' and '2/9' count as the same footer."""
    return re.sub(r"\d+", "#", line.strip().lower())


def strip_repeated_lines(pages):
    """Remove running headers/footers: lines that repeat on many pages.

    Printing a web page to PDF adds a date/title header and a URL + page
    counter footer to every page. They are not contract text, so drop them.
    """
    if len(pages) < 3:
        return pages
    counts = {}
    for lines in pages:
        for shape in {_shape(ln) for ln in lines if ln.strip()}:
            counts[shape] = counts.get(shape, 0) + 1
    limit = max(3, int(0.4 * len(pages)))
    return [[ln for ln in lines if counts.get(_shape(ln), 0) < limit or not ln.strip()]
            for lines in pages]


# ------------------------------------------------------------- cleaning -----
# A new clause usually starts with "1.", "2.3", "(a)", "Section 4", "ARTICLE V"
_NEW_CLAUSE = re.compile(
    r"^\s*(\d+(\.\d+)*[.)]?\s|\([a-zA-Z0-9]{1,4}\)\s|section\s+\d+|article\s+[ivx\d]+)",
    re.IGNORECASE,
)
_PAGE_NUMBER = re.compile(
    r"^\s*(page\s+\d+(\s+of\s+\d+)?|-?\s*\d{1,3}\s*-?|\d{1,3}\s*/\s*\d{1,3})\s*$",
    re.IGNORECASE,
)
# browser print leftovers: "10/4/26, 12:10 PM ..." and lines that are just a web address
_BROWSER_NOISE = re.compile(
    r"^\s*(\d{1,2}/\d{1,2}/\d{2,4},?\s+\d{1,2}:\d{2}\s*[AP]M\b.*|(https?://|www\.)\S+(\s+\d{1,3}\s*/\s*\d{1,3})?)\s*$",
    re.IGNORECASE,
)


def clean_text(raw):
    """Remove page numbers, hyphenation and line-wrap formatting."""
    text = unicodedata.normalize("NFKC", raw)
    text = text.replace("\u00ad", "")              # soft hyphens
    text = re.sub(r"-\n(?=[a-z])", "", text)       # "exam-\nple" -> "example"
    lines = [ln for ln in text.splitlines()
             if not _PAGE_NUMBER.match(ln) and not _BROWSER_NOISE.match(ln)]
    return "\n".join(lines)


def split_into_clauses(raw):
    """Clean the text and cut it into clause-sized pieces."""
    paragraphs, current = [], []
    for line in clean_text(raw).splitlines():
        stripped = line.strip()
        if not stripped:                            # blank line = paragraph break
            if current:
                paragraphs.append(" ".join(current))
                current = []
            continue
        if _NEW_CLAUSE.match(stripped) and current:  # numbered start = new clause
            paragraphs.append(" ".join(current))
            current = []
        current.append(stripped)
    if current:
        paragraphs.append(" ".join(current))

    pieces = []
    for para in paragraphs:
        para = re.sub(r"\s+", " ", para).strip()
        if len(para) <= MAX_CHARS:
            pieces.append(para)
            continue
        # very long paragraph: split by sentence, regroup up to MAX_CHARS
        sentences = re.split(r"(?<=[.;])\s+(?=[A-Z(])", para)
        buf = ""
        for s in sentences:
            if buf and len(buf) + len(s) + 1 > MAX_CHARS:
                pieces.append(buf)
                buf = s
            else:
                buf = f"{buf} {s}".strip()
        if buf:
            pieces.append(buf)

    # merge very short pieces (headings) into the next piece, drop leftovers
    clauses, carry = [], ""
    for p in pieces:
        p = f"{carry} {p}".strip()
        if len(p) < MIN_CHARS:
            carry = p
        else:
            clauses.append(p)
            carry = ""
    return clauses


# -------------------------------------------------------------- scoring -----
def score_clauses(vectorizer, model, clauses):
    """Predict risk, confidence, action and key words for each clause."""
    if not clauses:
        return pd.DataFrame(
            columns=["clause", "predicted_risk", "confidence", "action", "key_words"]
        )
    X = vectorizer.transform(clauses)
    proba = model.predict_proba(X)
    class_idx = proba.argmax(axis=1)
    classes = list(model.classes_)
    names = vectorizer.get_feature_names_out()

    rows = []
    for i, text in enumerate(clauses):
        risk = classes[class_idx[i]]
        conf = float(proba[i, class_idx[i]])
        row = X[i]
        coeffs = model.coef_[class_idx[i]]
        scored = sorted(
            ((names[j], coeffs[j]) for j in row.indices if coeffs[j] > 0),
            key=lambda t: t[1], reverse=True,
        )[:5]
        action = review_action(risk, conf)
        rows.append({
            "clause": text,
            "predicted_risk": risk,
            "confidence": round(conf, 3),
            "action": f"{ACTION_ICONS[action]} {action}",
            "key_words": ", ".join(w for w, _ in scored),
        })
    df = pd.DataFrame(rows)
    df["_order"] = df["action"].map(lambda a: ACTION_ORDER[a.split(" ", 1)[1]])
    return (df.sort_values(["_order", "confidence"])
              .drop(columns="_order").reset_index(drop=True))
